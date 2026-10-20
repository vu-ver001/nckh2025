"""
Tạo file .docx chứa tất cả các bảng kết quả từ báo cáo tổng kết đề tài.
Chạy: venv\\Scripts\\python.exe scripts/create_tables_docx.py
"""

from docx import Document
from docx.shared import Pt, Inches, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
import os


def set_cell_shading(cell, color_hex):
    """Set background color for a table cell."""
    shading = cell._element.get_or_add_tcPr()
    shd = shading.makeelement(qn("w:shd"), {
        qn("w:fill"): color_hex,
        qn("w:val"): "clear",
    })
    shading.append(shd)


def set_cell_border(cell, **kwargs):
    """Set cell border. kwargs: top, bottom, left, right, insideH, insideV"""
    tc = cell._element
    tcPr = tc.get_or_add_tcPr()
    tcBorders = tcPr.makeelement(qn("w:tcBorders"), {})
    for edge, val in kwargs.items():
        element = tcBorders.makeelement(qn(f"w:{edge}"), {
            qn("w:val"): val.get("val", "single"),
            qn("w:sz"): val.get("sz", "4"),
            qn("w:color"): val.get("color", "000000"),
            qn("w:space"): val.get("space", "0"),
        })
        tcBorders.append(element)
    tcPr.append(tcBorders)


def add_formatted_table(doc, title, headers, rows, caption=None):
    """Add a formatted table to the document."""
    # Title
    p = doc.add_paragraph()
    run = p.add_run(title)
    run.bold = True
    run.font.size = Pt(13)
    run.font.name = "Times New Roman"
    p.space_after = Pt(6)

    # Create table
    num_cols = len(headers)
    num_rows = len(rows) + 1
    table = doc.add_table(rows=num_rows, cols=num_cols)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"

    # Set column widths evenly
    for col_idx in range(num_cols):
        for row in table.rows:
            row.cells[col_idx].width = Cm(16.0 / num_cols)

    # Header row
    for i, header in enumerate(headers):
        cell = table.rows[0].cells[i]
        cell.text = ""
        p = cell.paragraphs[0]
        run = p.add_run(header)
        run.bold = True
        run.font.size = Pt(11)
        run.font.name = "Times New Roman"
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        set_cell_shading(cell, "D9E2F3")

    # Data rows
    for r_idx, row_data in enumerate(rows):
        for c_idx, cell_text in enumerate(row_data):
            cell = table.rows[r_idx + 1].cells[c_idx]
            cell.text = ""
            p = cell.paragraphs[0]
            run = p.add_run(str(cell_text))
            run.font.size = Pt(11)
            run.font.name = "Times New Roman"
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT

    # Caption
    if caption:
        p = doc.add_paragraph()
        run = p.add_run(caption)
        run.italic = True
        run.font.size = Pt(10)
        run.font.name = "Times New Roman"
        run.font.color.rgb = RGBColor(100, 100, 100)
        p.space_after = Pt(12)

    return table


def main():
    doc = Document()

    # --- Set default font ---
    style = doc.styles["Normal"]
    font = style.font
    font.name = "Times New Roman"
    font.size = Pt(12)

    # --- Set narrow margins ---
    for section in doc.sections:
        section.top_margin = Cm(2.0)
        section.bottom_margin = Cm(2.0)
        section.left_margin = Cm(2.5)
        section.right_margin = Cm(2.5)

    # ========== TITLE ==========
    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run(
        "BÁO CÁO TỔNG KẾT ĐỀ TÀI NGHIÊN CỨU KHOA HỌC SINH VIÊN\n"
        "CÁC BẢNG KẾT QUẢ VÀ MINH CHỨNG"
    )
    run.bold = True
    run.font.size = Pt(16)
    run.font.name = "Times New Roman"
    title.space_after = Pt(4)

    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = subtitle.add_run(
        "Nghiên cứu ứng dụng thuật toán YOLO và framework MediaPipe\n"
        "trong hệ thống giám sát phát hiện gian lận thi cử\n"
        "Ngày lập: 08/09/2026"
    )
    run.font.size = Pt(12)
    run.font.name = "Times New Roman"
    run.italic = True
    subtitle.space_after = Pt(18)

    # ========== BẢNG 1: Chương 2 — Mục tiêu cụ thể ==========
    add_formatted_table(
        doc,
        "Bảng 1. Mục tiêu cụ thể và tiêu chí kiểm chứng (Chương 2, §2.2)",
        ["Mục tiêu", "Tiêu chí kiểm chứng", "Trạng thái 08/09/2026"],
        [
            ["Nạp checkpoint Colab",
             "Model nạp được, ghi tên/đường dẫn/checksum, không âm thầm fallback",
             "Đạt trong MC02, MC07; cần xác nhận dashboard phiên cuối"],
            ["Định lượng detector",
             "Cùng 214 ảnh validation, confidence 0,30, IoU 0,50; có TP/FP/FN/P/R/F1",
             "Đạt; MC03 và MC05"],
            ["Nạp Pose + Hands",
             "Có model thật, version và số landmark",
             "Đạt kiểm tra kỹ thuật; MC06"],
            ["Luồng tích hợp",
             "Video → cảnh báo → snapshot → dashboard → dừng phiên",
             "Đạt bằng fixture; MC07"],
            ["Kiểm thử phần mềm",
             "Tất cả test hiện có và lỗi model/snapshot được kiểm tra",
             "57/57 đạt; MC08"],
            ["Đánh giá thực tế",
             "Sáu video độc lập, mốc sự kiện, FP/FN, latency, FPS",
             "Chưa bắt đầu vì thiếu video/nhãn; MC10"],
            ["Tuyên bố sẵn sàng phòng thi",
             "Có đánh giá độc lập và người duyệt xác nhận",
             "Chưa đạt, không được ghi là đã nghiệm thu"],
        ],
        "Nguồn: MC01–MC10 trong workspace; trạng thái tại 08/09/2026."
    )

    # ========== BẢNG 2: Chương 4 — Yêu cầu chức năng ==========
    add_formatted_table(
        doc,
        "Bảng 2. Yêu cầu chức năng (Chương 4, §4.1)",
        ["Mã", "Yêu cầu", "Cách kiểm tra", "Trạng thái"],
        [
            ["FR-01", "Chọn nguồn video/camera và bắt đầu phiên", "UI/API start; test hiện có", "Đã kiểm tra phần mềm"],
            ["FR-02", "Nạp đúng checkpoint Colab và hiển thị tên model", "status model + log", "Đã kiểm tra MC07; cần xác nhận dashboard"],
            ["FR-03", "Phát hiện box điện thoại", "detector/validation", "Đã đo trên 214 ảnh"],
            ["FR-04", "Trích xuất pose/hand landmarks", "MediaPipe thật", "Đã kiểm tra MC06/MC07"],
            ["FR-05", "Gộp tín hiệu và chống lặp cảnh báo", "state/temporal confirmation", "Đã có test; cần video độc lập"],
            ["FR-06", "Chụp snapshot kèm sự kiện", "endpoint snapshot", "HTTP 200 trong MC07"],
            ["FR-07", "Hiển thị trạng thái trên dashboard", "status, level, model, source", "Có pipeline; cần kiểm tra màn hình cuối"],
            ["FR-08", "Dừng phiên và giải phóng tài nguyên", "stop API", "Đạt trong MC07"],
        ],
        "Nguồn: Kiểm tra phần mềm, MC06, MC07."
    )

    # ========== BẢNG 3: Chương 5 — Thiết kế thực nghiệm ==========
    add_formatted_table(
        doc,
        "Bảng 3. Thiết kế thực nghiệm (Chương 5, §5.1)",
        ["Thành phần", "Điều kiện đã khóa"],
        [
            ["Tập detector", "214 ảnh validation từ phone_detection_v3"],
            ["Confidence", "0,30"],
            ["NMS IoU", "0,45"],
            ["IoU ghép nhãn", "0,50"],
            ["Checkpoint chính", "phone_v3_colab/best.pt, SHA-256 nêu tại §3.3"],
            ["Checkpoint đối chứng", "phone_v2 theo artifact comparison"],
        ],
        "Nguồn: scripts/compare_phone_v3.py; điều kiện được khóa trước khi đo."
    )

    # ========== BẢNG 4: Chương 5 — So sánh checkpoint ==========
    add_formatted_table(
        doc,
        "Bảng 4. So sánh checkpoint phone_v2 và phone_v3_colab (Chương 5, §5.3)",
        ["Checkpoint", "TP", "FP", "FN", "Precision", "Recall", "F1"],
        [
            ["phone_v2",    "41",  "137", "221", "23,03%", "15,65%", "18,64%"],
            ["phone_v3_colab", "124", "78",  "138", "61,39%", "47,33%", "53,45%"],
        ],
        "Nguồn: artifacts/phone-colab-comparison.json. Validation đã dùng chọn checkpoint, không phải test phòng thi."
    )

    # ========== BẢNG 5: Chương 5 — Đối chứng logic ==========
    add_formatted_table(
        doc,
        "Bảng 5. Đối chứng logic detector/fusion — phiên bản fixed (Chương 5, §5.4)",
        ["Biến thể fixed", "TP", "FP", "FN", "Precision", "Recall", "F1"],
        [
            ["YOLO-World",    "112", "69",  "150", "61,88%", "42,75%", "50,56%"],
            ["phone_v3_colab", "124", "73",  "138", "62,94%", "47,33%", "54,03%"],
            ["combined",      "124", "73",  "138", "62,94%", "47,33%", "54,03%"],
        ],
        "Nguồn: artifacts/research/detector-variants-validation-fixed.json. Tắt pose và phủ quyết vật được phép."
    )

    # ========== BẢNG 6: Chương 5 — Phân tích lỗi ==========
    add_formatted_table(
        doc,
        "Bảng 6. Phân tích lỗi theo kích thước điện thoại (Chương 5, §5.5)",
        ["Tỷ lệ diện tích GT/ảnh", "GT", "Detect", "Miss", "Recall"],
        [
            ["< 0,25%",  "91", "22", "69", "24,18%"],
            ["0,25% – 1%", "74", "39", "35", "52,70%"],
            ["≥ 1%",      "97", "63", "34", "64,95%"],
        ],
        "Nguồn: artifacts/research/error-analysis/summary.json. Nhóm diện tích do dự án định nghĩa, không phải COCO small/medium/large."
    )

    # ========== BẢNG 7: Chương 5 — Đánh giá video ==========
    add_formatted_table(
        doc,
        "Bảng 7. Đánh giá video độc lập — trạng thái chờ dữ liệu (Chương 5, §5.9)",
        ["Clip", "Nội dung dự kiến", "File", "Mốc sự kiện", "Kết quả"],
        [
            ["normal_01", "Ngồi làm bài, viết, nhìn giấy",     "Chưa có", "Chưa có", "Chưa đo"],
            ["normal_02", "Ngồi làm bài, viết, nhìn giấy",     "Chưa có", "Chưa có", "Chưa đo"],
            ["objects_01", "Máy tính/hộp bút/sổ",            "Chưa có", "Chưa có", "Chưa đo"],
            ["objects_02", "Máy tính/hộp bút/sổ",            "Chưa có", "Chưa có", "Chưa đo"],
            ["phone_01",  "Cầm điện thoại rõ trước camera",   "Chưa có", "Chưa có", "Chưa đo"],
            ["phone_02",  "Đặt xa hoặc che một phần",        "Chưa có", "Chưa có", "Chưa đo"],
        ],
        "Nguồn: demo/recordings/manifest.json. Lần chạy evaluator ngày 08/09: status=missing_video_or_labels, 0 clip sẵn sàng."
    )

    # ========== BẢNG 8: Kết luận theo mục tiêu ==========
    add_formatted_table(
        doc,
        "Bảng 8. Kết luận theo mục tiêu (Kết luận, §6.1)",
        ["Mục tiêu", "Kết luận có thể bảo vệ"],
        [
            ["Xây dựng nguyên mẫu YOLO + MediaPipe",
             "Đạt ở mức nguyên mẫu kỹ thuật; detector, landmarks, fusion, snapshot và dashboard đã nối được"],
            ["So sánh model Colab",
             "Đạt trên 214 ảnh validation; Colab tốt hơn phone_v2 theo các bảng có điều kiện"],
            ["Đạt chất lượng trên 90%",
             "Không đạt / không đủ dữ liệu để kết luận; số liệu gốc bị rút"],
            ["Đánh giá hành vi",
             "Mới có module và quy tắc; chưa có metric hành vi độc lập"],
            ["Đánh giá video thực tế",
             "Chưa thực hiện vì thiếu sáu video và nhãn"],
            ["Sẵn sàng phòng thi",
             "Chưa đủ căn cứ; chỉ nên trình diễn kỹ thuật có người giám sát"],
        ],
        None
    )

    # ========== BẢNG 9: Phụ lục 1 — Ma trận minh chứng ==========
    add_formatted_table(
        doc,
        "Bảng 9. Ma trận kết quả – minh chứng (Phụ lục 1)",
        ["Mã", "Tuyên bố được phép dùng", "Minh chứng", "Cách kiểm tra lại", "Điểm còn thiếu"],
        [
            ["R01", "Dataset 1.600 train/214 val, 2.177/262 box",
             "MC01", "Đọc JSON/manifest", "Quyền dùng và đồng nhất với Colab"],
            ["R02", "Checkpoint Colab là model một lớp phone, checksum nêu trên",
             "MC02", "SHA-256 + nạp model", "Xác nhận GPU/biên nhận Colab"],
            ["R03", "Colab vượt phone_v2 ở bảng raw",
             "MC03", "Chạy lại cùng validation", "Phân biệt raw/fixed khi trích hồ sơ"],
            ["R04", "Fixed phone_v3 đạt P62,94 / R47,33 / F1 54,03",
             "MC05", "Kiểm tra JSON và script version", "Chưa phải test ngoài mẫu"],
            ["R05", "Điện thoại nhỏ là nhóm bỏ sót nhiều",
             "MC04", "Đọc bảng size", "Nguyên nhân semantic cần review người"],
            ["R06", "MediaPipe Pose + Hands nạp thật",
             "MC06", "Mở summary/overlay", "Chưa có metric hành vi"],
            ["R07", "Luồng tích hợp tạo event/snapshot và stop được",
             "MC07", "Mở summary/snapshot", "Fixture lặp ảnh validation"],
            ["R08", "Bộ test 57/57, ổn định thành phần 1.800,259 s",
             "MC08", "Đọc log/JSON", "Chưa đại diện camera/WebSocket"],
            ["R09", "Có 280 ứng viên cần review",
             "MC09", "Đọc manifest", "Chưa dùng làm negative train"],
            ["R10", "Sáu video độc lập chưa sẵn sàng",
             "MC10", "events=null, thiếu file", "Cần quay, gán nhãn, chạy đánh giá"],
        ],
        "Nguồn: các artifact trong workspace. Chỉ dùng kết quả đã được ghi rõ điều kiện đo."
    )

    # ========== BẢNG 10: Phụ lục 5 — Việc còn thiếu ==========
    add_formatted_table(
        doc,
        "Bảng 10. Việc còn thiếu và người xác nhận (Phụ lục 5)",
        ["Ưu tiên", "Việc", "Người cần xác nhận", "Bằng chứng đầu ra"],
        [
            ["P0", "Tình trạng nghiệm thu, mã đề tài, bản chính thức",
             "GVHD/Khoa", "quyết định/biên bản"],
            ["P0", "Sáu video + mốc nhãn",
             "Nhóm + người review", "video, manifest, label"],
            ["P0", "Số liệu nào được phép đưa vào hồ sơ",
             "GVHD", "bảng duyệt số liệu"],
            ["P1", "Quyền dữ liệu, đạo đức, sử dụng AI",
             "Nhóm + Khoa", "phiếu xác nhận"],
            ["P1", "Review ứng viên negative",
             "Nhóm", "manifest corrected"],
            ["P2", "Train tiếp nếu lỗi model",
             "Nhóm trên Colab", "notebook, checkpoint, checksum"],
        ],
        "Điểm dừng hồ sơ: chỉ sau khi P0 hoàn tất mới có thể thay \"chưa đo/chưa xác nhận\" bằng kết quả."
    )

    # ========== BẢNG 11: Đính chính từ báo cáo gốc ==========
    add_formatted_table(
        doc,
        "Bảng 11. Đính chính các tuyên bố từ báo cáo gốc (Phụ lục DC)",
        ["Mã", "Tuyên bố bản gốc", "Bằng chứng hiện có", "Quyết định nội dung"],
        [
            ["DC01", "Accuracy 91,2%; F1 hành vi 85,3%",
             "Chưa có log, nhãn video/hành vi hoặc phép đo tương ứng trong workspace",
             "RÚT khỏi kết quả đã đạt"],
            ["DC02", "YOLOv8m, 100 epoch, khoảng 3,5 giờ, RTX 3060",
             "args.yaml ghi YOLOv8n, 40 epoch, batch 16, imgsz 640; GPU chưa xác minh",
             "SỬA mô tả lượt Colab theo args/checkpoint"],
            ["DC03", "Khoảng 2.690 ảnh gốc và 8.990 ảnh sau augmentation",
             "1.600 train + 214 validation từ CITW/COCO, 2.177/262 box",
             "SỬA theo manifest và audit"],
            ["DC04", "MediaPipe đã được dùng trong phiên bản báo cáo ban đầu",
             "Mã cũ là YOLO-Pose; MediaPipe tích hợp 06/09/2026",
             "TÁCH MỐC THỜI GIAN"],
            ["DC05", "YOLO đạt ~93,4% precision, 92,1% recall, 93,3% mAP",
             "Colab P61,39%, R47,33%, F1 53,45% ở conf 0,30, IoU 0,50",
             "RÚT số cũ; giữ bảng mới với đầy đủ điều kiện"],
            ["DC06", "Hệ thống 52 FPS, 19 ms, ổn định 90 phút",
             "3,59 FPS trên fixture; 1.800,259 s/6.460 frame; chưa có benchmark độc lập",
             "RÚT hoặc ghi chưa xác minh"],
            ["DC07", "Kiến trúc lai chứng minh hiệu quả hơn",
             "Combined = phone độc lập khi tắt phủ quyết; World không cải thiện phone",
             "HẠN CHẾ KẾT LUẬN"],
            ["DC08", "Hệ thống đã kiểm thử thực tế/phòng thi",
             "Kiểm thử kỹ thuật với ảnh validation lặp; chưa có video nhóm",
             "SỬA thành kiểm thử kỹ thuật"],
            ["DC09", "Chưa có công bố khoa học",
             "Workspace không có bài báo, DOI hoặc thư chấp nhận",
             "GIỮ với điều kiện nhóm/GVHD xác nhận"],
        ],
        "Nguồn: outputs/01_doi_chieu_va_dinh_chinh.md; đây là bản làm việc, không phải quyết định nghiệm thu."
    )

    # ========== LƯU FILE ==========
    output_dir = os.path.join(os.path.dirname(__file__), "..", "docs")
    os.makedirs(output_dir, exist_ok=True)
    output_path = os.path.join(output_dir, "BANG_BAO_CAO_TONG_KET.docx")
    doc.save(output_path)
    print(f"Da tao file: {output_path}")
    print(f"Kich thuoc: {os.path.getsize(output_path):,} bytes")
    print("Hoan tat!")


if __name__ == "__main__":
    main()
