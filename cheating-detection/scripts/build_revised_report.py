"""Build the review DOCX from the evidence-based Markdown using bundled python-docx."""
from pathlib import Path
import re
import argparse
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from PIL import Image

ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT/'docs/REPORT_REVISED_2026-09-06.md'
DESTINATION = ROOT/'deliverables/Bao_cao_sua_de_GVHD_ra_soat_2026-09-06.docx'


def plain(text):
    return text.replace('**','').replace('*','').replace('`','')


def main(source=SOURCE, destination=DESTINATION):
    document = Document()
    document.settings.odd_and_even_pages_header_footer = False
    section = document.sections[0]
    section.different_first_page_header_footer = False
    section.page_width, section.page_height = Inches(8.5), Inches(11)
    section.top_margin = section.bottom_margin = Inches(.95)
    section.header_distance = section.footer_distance = Inches(.45)
    section.left_margin = section.right_margin = Inches(.85)
    normal = document.styles['Normal']
    normal.font.name, normal.font.size = 'Times New Roman', Pt(12)
    normal.paragraph_format.line_spacing = 1.12
    normal.paragraph_format.space_after = Pt(7)
    for name, size in [('Title',18),('Heading 1',15),('Heading 2',13)]:
        style = document.styles[name]
        style.font.name, style.font.size = 'Times New Roman', Pt(size)
        style.font.color.rgb = RGBColor(0,0,0)
        style.font.bold = True
        style.paragraph_format.space_before = Pt(14)
        style.paragraph_format.space_after = Pt(8)
        style.paragraph_format.keep_with_next = True
    # Remove inherited theme fonts and paragraph borders, including the default
    # Title rule, so the Word rendering matches the specified plain typography.
    for name in ('Normal','Title','Heading 1','Heading 2'):
        style = document.styles[name]
        fonts = style.element.get_or_add_rPr().get_or_add_rFonts()
        for attr in list(fonts.attrib):
            if attr.endswith('Theme'):
                del fonts.attrib[attr]
        for attr in ('ascii','hAnsi','eastAsia','cs'):
            fonts.set(qn('w:'+attr),'Times New Roman')
    for border in list(document.styles.element.iter(qn('w:pBdr'))):
        border.getparent().remove(border)
    header = section.header.paragraphs[0]
    header.text = 'HUNRE  •  BẢN SỬA PHỤC VỤ RÀ SOÁT'
    header.runs[0].font.size = Pt(9)
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    footer.add_run('Chưa thay thế hồ sơ đã nộp  |  Trang ').font.size = Pt(9)
    field = OxmlElement('w:fldSimple')
    field.set(qn('w:instr'), 'PAGE')
    footer._p.append(field)
    source = Path(source)
    destination = Path(destination)
    lines = source.read_text(encoding='utf-8').splitlines()
    index = 0
    while index < len(lines):
        line = lines[index].strip()
        if not line:
            index += 1
            continue
        if line.startswith('# '):
            paragraph = document.add_paragraph(plain(line[2:]), 'Title')
            paragraph.paragraph_format.space_before = Pt(40)
        elif line.startswith('## '):
            title = plain(line[3:])
            paragraph = document.add_paragraph(title, 'Heading 1')
            if title in ('Tóm tắt','Phụ lục nội dung cần đính chính và bằng chứng'):
                paragraph.paragraph_format.page_break_before = True
        elif line.startswith('### '):
            document.add_paragraph(plain(line[4:]), 'Heading 2')
        elif line.startswith('!['):
            match = re.match(r'!\[(.*?)\]\((.*?)\)',line)
            path = (source.parent/match.group(2)).resolve()
            with Image.open(path) as picture:
                width, height = picture.size
            picture_width = min(6.6, 4.0*width/height)
            paragraph = document.add_paragraph()
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            paragraph.paragraph_format.keep_with_next = True
            paragraph.add_run().add_picture(str(path),width=Inches(picture_width))
        elif line.startswith('|'):
            rows=[]
            while index < len(lines) and lines[index].strip().startswith('|'):
                cells=[cell.strip() for cell in lines[index].strip().strip('|').split('|')]
                if not all(re.fullmatch(r'[-: ]+',cell) for cell in cells):
                    rows.append(cells)
                index += 1
            table=document.add_table(rows=0, cols=len(rows[0]))
            table.alignment=WD_TABLE_ALIGNMENT.CENTER
            table.autofit=False
            widths = {2:[2.0,4.8],5:[2.6,1.0,1.1,1.,1.1],7:[1.5,.6,.6,.6,1.2,1.15,1.15]}.get(len(rows[0]),[6.8/len(rows[0])]*len(rows[0]))
            for column,width in zip(table.columns,widths):
                column.width=Inches(width)
            for ri,values in enumerate(rows):
                row=table.add_row()
                cannot_split=OxmlElement('w:cantSplit')
                row._tr.get_or_add_trPr().append(cannot_split)
                if ri==0:
                    repeat=OxmlElement('w:tblHeader')
                    row._tr.get_or_add_trPr().append(repeat)
                for ci,value in enumerate(values):
                    cell=row.cells[ci]
                    cell.width=Inches(widths[ci])
                    cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
                    cell.text=plain(value)
                    props=cell._tc.get_or_add_tcPr()
                    borders=OxmlElement('w:tcBorders')
                    for side in ('top','left','bottom','right'):
                        border=OxmlElement('w:'+side)
                        for key,val in [('val','single'),('sz','4'),('color','D9D9D9')]:
                            border.set(qn('w:'+key),val)
                        borders.append(border)
                    props.append(borders)
                    margins=OxmlElement('w:tcMar')
                    for side in ('top','left','bottom','right'):
                        margin=OxmlElement('w:'+side)
                        margin.set(qn('w:w'),'90')
                        margin.set(qn('w:type'),'dxa')
                        margins.append(margin)
                    props.append(margins)
                    shade=OxmlElement('w:shd')
                    shade.set(qn('w:fill'),'E8EEF4' if ri==0 else 'FFFFFF')
                    props.append(shade)
                    for paragraph in cell.paragraphs:
                        paragraph.paragraph_format.space_after=Pt(3)
                        paragraph.paragraph_format.line_spacing=1.05
                        if len(rows)<=6 and ri<len(rows)-1:
                            paragraph.paragraph_format.keep_with_next=True
                        if ci>0 and len(widths)>=5:
                            paragraph.alignment=WD_ALIGN_PARAGRAPH.CENTER
                        for run in paragraph.runs:
                            run.font.size=Pt(11)
                            run.bold=ri==0
            document.add_paragraph().paragraph_format.space_after=Pt(3)
            continue
        else:
            if line.startswith('> '):
                line = line[2:]
            if line.startswith('- '):
                paragraph = document.add_paragraph(plain(line[2:]), 'List Bullet')
            else:
                paragraph=document.add_paragraph(plain(line))
            if line.startswith('Hình '):
                paragraph.alignment=WD_ALIGN_PARAGRAPH.CENTER
                for run in paragraph.runs:
                    run.font.size=Pt(10)
                    run.italic=True
        index += 1
    destination.parent.mkdir(parents=True, exist_ok=True)
    document.core_properties.title='Bản sửa báo cáo nghiên cứu YOLO và MediaPipe'
    document.core_properties.subject='Bản làm việc theo bằng chứng để nhóm và GVHD rà soát'
    document.save(destination)
    print(destination)


if __name__=='__main__':
    parser = argparse.ArgumentParser(description='Build a review DOCX from Markdown.')
    parser.add_argument('--source', default=str(SOURCE), help='Markdown source path')
    parser.add_argument('--destination', default=str(DESTINATION), help='DOCX output path')
    arguments = parser.parse_args()
    main(arguments.source, arguments.destination)
